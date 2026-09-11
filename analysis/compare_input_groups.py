"""Compare citation behaviour in real- and synthetic-input Hercules judgments.

Inputs
------
data/processed/judgment_citations_59.csv
data/processed/unique_citation_review.xlsx

Outputs
-------
data/processed/input_group_comparison_per_judgment.csv
data/processed/input_group_comparison_per_input_record.csv
data/processed/input_group_statistical_comparison.csv
data/processed/input_group_statistical_comparison.json

The inferential results are exploratory. Repeated generations, re-uploaded cases,
and convenience sampling mean that the observations are not guaranteed to be
independent or representative of all Hercules inputs.
"""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from statistics import mean, median, stdev
import csv
import json
import math
import random
import re

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PROCESSED = PROJECT_ROOT / "data" / "processed"

CITATIONS_FILE = PROCESSED / "judgment_citations_59.csv"
REVIEW_FILE = PROCESSED / "unique_citation_review.xlsx"

PER_JUDGMENT_OUTPUT = (
    PROCESSED / "input_group_comparison_per_judgment.csv"
)
PER_INPUT_OUTPUT = (
    PROCESSED / "input_group_comparison_per_input_record.csv"
)
STATISTICS_CSV_OUTPUT = (
    PROCESSED / "input_group_statistical_comparison.csv"
)
STATISTICS_JSON_OUTPUT = (
    PROCESSED / "input_group_statistical_comparison.json"
)

RANDOM_SEED = 748
BOOTSTRAP_ITERATIONS = 20_000
PERMUTATION_ITERATIONS = 50_000
EXPECTED_JUDGMENTS = 59
EXPECTED_CITATION_RECORDS = 115


def clean(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_authority(value: object) -> str:
    text = clean(value).casefold()
    text = text.replace("&", " and ")
    text = text.replace("’", "'")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def read_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        raise FileNotFoundError(f"Required file not found: {path}")

    with path.open("r", encoding="utf-8-sig", newline="") as file:
        return list(csv.DictReader(file))


def load_review_mapping() -> dict[str, str]:
    if not REVIEW_FILE.exists():
        raise FileNotFoundError(f"Required file not found: {REVIEW_FILE}")

    workbook = load_workbook(REVIEW_FILE, data_only=True, read_only=True)

    if "unique_citation_review" not in workbook.sheetnames:
        raise ValueError(
            "Worksheet 'unique_citation_review' was not found in "
            f"{REVIEW_FILE}"
        )

    worksheet = workbook["unique_citation_review"]
    headers = [clean(cell.value) for cell in worksheet[1]]
    index = {header: position for position, header in enumerate(headers)}

    required = {
        "authority",
        "citation_variants",
        "citation_correct_external",
    }
    missing = required - set(index)

    if missing:
        raise ValueError(
            "Review workbook is missing columns: "
            + ", ".join(sorted(missing))
        )

    mapping: dict[str, str] = {}

    for values in worksheet.iter_rows(min_row=2, values_only=True):
        authority = clean(values[index["authority"]])
        status = clean(
            values[index["citation_correct_external"]]
        ).lower()

        if not authority:
            continue

        if status not in {"yes", "incomplete", "no"}:
            raise ValueError(
                f"Unexpected citation review status for {authority}: "
                f"{status!r}"
            )

        variants = [authority]
        variants.extend(
            part.strip()
            for part in clean(
                values[index["citation_variants"]]
            ).split("||")
            if part.strip()
        )

        for variant in variants:
            key = normalize_authority(variant)

            if not key:
                continue

            previous = mapping.get(key)
            if previous is not None and previous != status:
                raise ValueError(
                    "Conflicting review labels for citation variant "
                    f"{variant!r}: {previous!r} and {status!r}"
                )

            mapping[key] = status

    return mapping


def classify_citation(
    citation_row: dict[str, str],
    mapping: dict[str, str],
) -> str:
    candidates = [
        citation_row.get("citation", ""),
        citation_row.get("citation_normalized", ""),
    ]

    for candidate in candidates:
        key = normalize_authority(candidate)
        if key in mapping:
            return mapping[key]

    raise ValueError(
        "Citation could not be matched to unique_citation_review.xlsx: "
        f"draft_id={citation_row.get('draft_id')!r}, "
        f"citation={citation_row.get('citation')!r}"
    )


def build_judgment_records(
    citation_rows: list[dict[str, str]],
    review_mapping: dict[str, str],
) -> list[dict[str, object]]:
    records: dict[str, dict[str, object]] = {}

    for citation in citation_rows:
        draft_id = clean(citation.get("draft_id"))
        input_type = clean(citation.get("input_type")).lower()
        case_doc_id = clean(citation.get("case_doc_id"))

        if not draft_id or not case_doc_id:
            raise ValueError(
                "Every citation record must contain draft_id and case_doc_id."
            )

        if input_type not in {"real", "synthetic"}:
            raise ValueError(
                f"Unexpected input_type for {draft_id}: {input_type!r}"
            )

        if draft_id not in records:
            records[draft_id] = {
                "draft_id": draft_id,
                "case_doc_id": case_doc_id,
                "input_type": input_type,
                "citation_mentions": 0,
                "correct_mentions": 0,
                "incomplete_mentions": 0,
                "incorrect_mentions": 0,
            }

        record = records[draft_id]

        if (
            record["input_type"] != input_type
            or record["case_doc_id"] != case_doc_id
        ):
            raise ValueError(
                f"Inconsistent metadata found for draft {draft_id}."
            )

        status = classify_citation(citation, review_mapping)
        record["citation_mentions"] += 1

        if status == "yes":
            record["correct_mentions"] += 1
        elif status == "incomplete":
            record["incomplete_mentions"] += 1
        else:
            record["incorrect_mentions"] += 1

    result: list[dict[str, object]] = []

    for record in records.values():
        total = int(record["citation_mentions"])
        incomplete = int(record["incomplete_mentions"])
        incorrect = int(record["incorrect_mentions"])

        record["strict_defect_mentions"] = incomplete + incorrect
        record["strict_defect_rate"] = (
            (incomplete + incorrect) / total if total else 0.0
        )
        record["incorrect_rate"] = incorrect / total if total else 0.0
        result.append(record)

    result.sort(
        key=lambda row: (
            str(row["input_type"]),
            str(row["case_doc_id"]),
            str(row["draft_id"]),
        )
    )
    return result


def build_input_records(
    judgment_records: list[dict[str, object]],
) -> list[dict[str, object]]:
    grouped: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)

    for record in judgment_records:
        key = (
            str(record["input_type"]),
            str(record["case_doc_id"]),
        )
        grouped[key].append(record)

    output: list[dict[str, object]] = []

    for (input_type, case_doc_id), records in grouped.items():
        judgment_count = len(records)
        total_mentions = sum(
            int(row["citation_mentions"]) for row in records
        )
        correct_mentions = sum(
            int(row["correct_mentions"]) for row in records
        )
        incomplete_mentions = sum(
            int(row["incomplete_mentions"]) for row in records
        )
        incorrect_mentions = sum(
            int(row["incorrect_mentions"]) for row in records
        )
        strict_defects = incomplete_mentions + incorrect_mentions

        output.append(
            {
                "case_doc_id": case_doc_id,
                "input_type": input_type,
                "judgment_count": judgment_count,
                "citation_mentions": total_mentions,
                "mean_citations_per_judgment": (
                    total_mentions / judgment_count
                ),
                "correct_mentions": correct_mentions,
                "incomplete_mentions": incomplete_mentions,
                "incorrect_mentions": incorrect_mentions,
                "strict_defect_rate": (
                    strict_defects / total_mentions
                    if total_mentions
                    else 0.0
                ),
                "incorrect_rate": (
                    incorrect_mentions / total_mentions
                    if total_mentions
                    else 0.0
                ),
            }
        )

    output.sort(
        key=lambda row: (
            str(row["input_type"]),
            str(row["case_doc_id"]),
        )
    )
    return output


def descriptive(values: list[float]) -> dict[str, float | int]:
    if not values:
        raise ValueError("Cannot summarise an empty group.")

    return {
        "n": len(values),
        "mean": mean(values),
        "median": median(values),
        "standard_deviation": stdev(values) if len(values) > 1 else 0.0,
        "minimum": min(values),
        "maximum": max(values),
    }


def hedges_g(real: list[float], synthetic: list[float]) -> float:
    n_real = len(real)
    n_synthetic = len(synthetic)

    if n_real < 2 or n_synthetic < 2:
        return 0.0

    variance_real = stdev(real) ** 2
    variance_synthetic = stdev(synthetic) ** 2
    degrees_freedom = n_real + n_synthetic - 2

    pooled_variance = (
        (n_real - 1) * variance_real
        + (n_synthetic - 1) * variance_synthetic
    ) / degrees_freedom

    if pooled_variance <= 0:
        return 0.0

    cohen_d = (
        mean(synthetic) - mean(real)
    ) / math.sqrt(pooled_variance)
    correction = 1 - (3 / (4 * degrees_freedom - 1))
    return correction * cohen_d


def bootstrap_difference_interval(
    real: list[float],
    synthetic: list[float],
    seed: int,
) -> tuple[float, float]:
    rng = random.Random(seed)
    differences: list[float] = []

    for _ in range(BOOTSTRAP_ITERATIONS):
        sampled_real = [rng.choice(real) for _ in real]
        sampled_synthetic = [
            rng.choice(synthetic) for _ in synthetic
        ]
        differences.append(
            mean(sampled_synthetic) - mean(sampled_real)
        )

    differences.sort()
    lower_index = int(0.025 * (len(differences) - 1))
    upper_index = int(0.975 * (len(differences) - 1))
    return differences[lower_index], differences[upper_index]


def permutation_p_value(
    real: list[float],
    synthetic: list[float],
    seed: int,
) -> float:
    rng = random.Random(seed)
    observed = abs(mean(synthetic) - mean(real))
    combined = real + synthetic
    real_size = len(real)
    exceedances = 0

    for _ in range(PERMUTATION_ITERATIONS):
        shuffled = combined.copy()
        rng.shuffle(shuffled)
        permuted_real = shuffled[:real_size]
        permuted_synthetic = shuffled[real_size:]
        difference = abs(
            mean(permuted_synthetic) - mean(permuted_real)
        )

        if difference >= observed - 1e-12:
            exceedances += 1

    return (exceedances + 1) / (PERMUTATION_ITERATIONS + 1)


def compare_metric(
    records: list[dict[str, object]],
    field: str,
    analysis_level: str,
    seed_offset: int,
) -> dict[str, object]:
    real = [
        float(record[field])
        for record in records
        if record["input_type"] == "real"
    ]
    synthetic = [
        float(record[field])
        for record in records
        if record["input_type"] == "synthetic"
    ]

    real_summary = descriptive(real)
    synthetic_summary = descriptive(synthetic)
    difference = mean(synthetic) - mean(real)
    interval_low, interval_high = bootstrap_difference_interval(
        real,
        synthetic,
        RANDOM_SEED + seed_offset,
    )

    return {
        "analysis_level": analysis_level,
        "metric": field,
        "real": real_summary,
        "synthetic": synthetic_summary,
        "mean_difference_synthetic_minus_real": difference,
        "bootstrap_95_ci_low": interval_low,
        "bootstrap_95_ci_high": interval_high,
        "permutation_p_value_two_sided": permutation_p_value(
            real,
            synthetic,
            RANDOM_SEED + 10_000 + seed_offset,
        ),
        "hedges_g": hedges_g(real, synthetic),
    }


def write_records(
    path: Path,
    records: list[dict[str, object]],
    columns: list[str],
) -> None:
    with path.open("w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=columns)
        writer.writeheader()

        for record in records:
            writer.writerow(
                {
                    column: record.get(column, "")
                    for column in columns
                }
            )


def main() -> None:
    citation_rows = read_csv(CITATIONS_FILE)

    if len(citation_rows) != EXPECTED_CITATION_RECORDS:
        raise ValueError(
            f"Expected {EXPECTED_CITATION_RECORDS} citation records, "
            f"found {len(citation_rows)}. Check the input file."
        )

    review_mapping = load_review_mapping()
    judgment_records = build_judgment_records(
        citation_rows,
        review_mapping,
    )

    if len(judgment_records) != EXPECTED_JUDGMENTS:
        raise ValueError(
            f"Expected {EXPECTED_JUDGMENTS} judgments, "
            f"found {len(judgment_records)}. Judgments with zero extracted "
            "citations must be added from the draft manifest."
        )

    input_records = build_input_records(judgment_records)

    comparisons: list[dict[str, object]] = []
    judgment_metrics = [
        "citation_mentions",
        "strict_defect_rate",
        "incorrect_rate",
    ]
    input_metrics = [
        "mean_citations_per_judgment",
        "strict_defect_rate",
        "incorrect_rate",
    ]

    for index, metric in enumerate(judgment_metrics):
        comparisons.append(
            compare_metric(
                judgment_records,
                metric,
                "judgment",
                index,
            )
        )

    for index, metric in enumerate(input_metrics, start=100):
        comparisons.append(
            compare_metric(
                input_records,
                metric,
                "input_record",
                index,
            )
        )

    write_records(
        PER_JUDGMENT_OUTPUT,
        judgment_records,
        [
            "draft_id",
            "case_doc_id",
            "input_type",
            "citation_mentions",
            "correct_mentions",
            "incomplete_mentions",
            "incorrect_mentions",
            "strict_defect_mentions",
            "strict_defect_rate",
            "incorrect_rate",
        ],
    )

    write_records(
        PER_INPUT_OUTPUT,
        input_records,
        [
            "case_doc_id",
            "input_type",
            "judgment_count",
            "citation_mentions",
            "mean_citations_per_judgment",
            "correct_mentions",
            "incomplete_mentions",
            "incorrect_mentions",
            "strict_defect_rate",
            "incorrect_rate",
        ],
    )

    flat_rows = []
    for comparison in comparisons:
        flat_rows.append(
            {
                "analysis_level": comparison["analysis_level"],
                "metric": comparison["metric"],
                "real_n": comparison["real"]["n"],
                "real_mean": comparison["real"]["mean"],
                "real_median": comparison["real"]["median"],
                "real_standard_deviation": comparison["real"][
                    "standard_deviation"
                ],
                "synthetic_n": comparison["synthetic"]["n"],
                "synthetic_mean": comparison["synthetic"]["mean"],
                "synthetic_median": comparison["synthetic"]["median"],
                "synthetic_standard_deviation": comparison[
                    "synthetic"
                ]["standard_deviation"],
                "mean_difference_synthetic_minus_real": comparison[
                    "mean_difference_synthetic_minus_real"
                ],
                "bootstrap_95_ci_low": comparison[
                    "bootstrap_95_ci_low"
                ],
                "bootstrap_95_ci_high": comparison[
                    "bootstrap_95_ci_high"
                ],
                "permutation_p_value_two_sided": comparison[
                    "permutation_p_value_two_sided"
                ],
                "hedges_g": comparison["hedges_g"],
            }
        )

    write_records(
        STATISTICS_CSV_OUTPUT,
        flat_rows,
        list(flat_rows[0].keys()),
    )

    report = {
        "method": {
            "comparison": "synthetic minus real",
            "bootstrap_iterations": BOOTSTRAP_ITERATIONS,
            "permutation_iterations": PERMUTATION_ITERATIONS,
            "random_seed": RANDOM_SEED,
            "strict_defect_definition": (
                "incomplete citation or incorrect citation"
            ),
            "review_label_source": str(REVIEW_FILE),
        },
        "sample": {
            "citation_records": len(citation_rows),
            "judgments": len(judgment_records),
            "input_records": len(input_records),
        },
        "comparisons": comparisons,
        "limitations": [
            "The tests are exploratory and do not establish causation.",
            "Repeated generations from the same input are not fully independent.",
            "Different document IDs may represent re-uploads of the same underlying case.",
            "Citation correctness labels were assigned at authority-group level and inherited by each occurrence.",
            "The convenience sample is not necessarily representative of all Hercules cases.",
        ],
    }

    with STATISTICS_JSON_OUTPUT.open("w", encoding="utf-8") as file:
        json.dump(report, file, indent=2)

    print("\nReal-versus-synthetic comparison completed")
    print("-------------------------------------------")
    print(f"Citation records: {len(citation_rows)}")
    print(f"Judgments: {len(judgment_records)}")
    print(f"Input records: {len(input_records)}")

    for comparison in comparisons:
        print(
            f"\n{comparison['analysis_level']} / "
            f"{comparison['metric']}"
        )
        print(
            f"  real mean: {comparison['real']['mean']:.4f} "
            f"(n={comparison['real']['n']})"
        )
        print(
            f"  synthetic mean: "
            f"{comparison['synthetic']['mean']:.4f} "
            f"(n={comparison['synthetic']['n']})"
        )
        print(
            "  difference (synthetic - real): "
            f"{comparison['mean_difference_synthetic_minus_real']:.4f}"
        )
        print(
            "  bootstrap 95% CI: "
            f"[{comparison['bootstrap_95_ci_low']:.4f}, "
            f"{comparison['bootstrap_95_ci_high']:.4f}]"
        )
        print(
            "  permutation p-value: "
            f"{comparison['permutation_p_value_two_sided']:.4f}"
        )
        print(f"  Hedges' g: {comparison['hedges_g']:.4f}")

    print(f"\nPer-judgment output: {PER_JUDGMENT_OUTPUT}")
    print(f"Per-input output: {PER_INPUT_OUTPUT}")
    print(f"Statistics CSV: {STATISTICS_CSV_OUTPUT}")
    print(f"Statistics JSON: {STATISTICS_JSON_OUTPUT}")
    print(
        "\nInterpret inferential results as exploratory because repeated "
        "generations and re-uploaded cases may not be independent."
    )


if __name__ == "__main__":
    main()
