"""
Pre-fill holding-accuracy and reasoning-grounding draft labels.

Sends the generated principle and retrieved evidence to GPT-4o-mini
and writes draft R1 labels into the annotation workbook for human
review and correction.

For holding accuracy: tries to fetch the authoritative judgment text.
If the source is blocked (e.g. BAILII bot protection), falls back to
LLM knowledge and flags the assessment accordingly.

For reasoning grounding: uses only the retrieved evidence already in
the workbook. No external fetch needed.

Usage:
    set OPENAI_API_KEY=sk-...
    python src/prefill_holding_grounding.py [--dry-run] [--limit N]
"""

import argparse
import json
import os
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook
from openai import OpenAI


PROJECT_ROOT = Path(__file__).resolve().parent.parent

WORKBOOK_PATH = (
    PROJECT_ROOT
    / "annotations"
    / "holding_grounding_annotation_115.xlsx"
)

CACHE_DIR = PROJECT_ROOT / "data" / "judgment_cache"

MODEL = "gpt-4o-mini"

HOLDING_SYSTEM_PROMPT_WITH_SOURCE = """\
You are a legal-annotation assistant. You will receive:

1. A CITATION to a legal authority.
2. A GENERATED PRINCIPLE that an AI system attributed to that authority.
3. The SOURCE JUDGMENT TEXT from the authoritative judgment.

Assess whether the generated principle accurately represents what \
the authority established.

Labels:
- yes: Materially faithful. Wording may be paraphrased but the legal \
meaning, scope, conditions and procedural context are preserved.
- partial: Some support, but the principle is incomplete, overstated, \
decontextualised or imprecise. A correct general idea is present but \
an important limitation, test, condition, or distinction between \
ratio and obiter is omitted.
- no: Unsupported, contradicted or attributed to the wrong authority. \
The authority concerns a different doctrine, reaches the opposite \
position, or does not establish the proposition.
- unclear: Insufficient evidence to decide reliably.
- n/a: Procedural rule or statutory provision; no case holding is \
being attributed.

An obiter observation is not automatically inaccurate. Label partial \
if presented as binding holding without qualification. Label no only \
if the supposed rule cannot reasonably be supported by the judgment.

Respond with valid JSON only:
{
  "holding_accuracy": "yes|partial|no|unclear|n/a",
  "source_paragraphs": "paragraph numbers or section references used",
  "holding_evidence": "2-3 sentence explanation",
  "confidence": "high|medium|low"
}
"""

HOLDING_SYSTEM_PROMPT_NO_SOURCE = """\
You are a legal-annotation assistant. You will receive:

1. A CITATION to a legal authority.
2. A GENERATED PRINCIPLE that an AI system attributed to that authority.

The source judgment text could not be fetched automatically. Use your \
knowledge of this legal authority to assess whether the generated \
principle accurately represents what the authority established.

Be conservative: if you are not confident about the actual holding, \
use "unclear" rather than guessing.

Labels:
- yes: Materially faithful based on established understanding of this \
authority.
- partial: Some support, but the principle is incomplete, overstated, \
decontextualised or imprecise.
- no: Unsupported, contradicted or attributed to the wrong authority.
- unclear: You are not confident enough about the actual holding to \
make a reliable assessment. DEFAULT TO THIS if uncertain.
- n/a: Procedural rule or statutory provision; no case holding is \
being attributed.

Respond with valid JSON only:
{
  "holding_accuracy": "yes|partial|no|unclear|n/a",
  "source_paragraphs": "not available - source not fetched",
  "holding_evidence": "2-3 sentence explanation; note that this is based on training knowledge, not source text",
  "confidence": "high|medium|low"
}
"""

GROUNDING_SYSTEM_PROMPT = """\
You are a legal-annotation assistant. You will receive:

1. A CITATION to a legal authority.
2. A GENERATED APPLICATION — how an AI system applied the authority \
in a draft judgment.
3. RETRIEVED PASSAGES — the actual passages the AI system retrieved \
before generating the draft.

Assess whether the retrieved passages support the generated \
application.

Important: check whether the RETRIEVAL supports the application, not \
whether the application is legally correct in general. An application \
may be legally sound but still ungrounded if the retrieved material \
does not contain the relevant proposition.

Labels:
- yes: Clear support. At least one retrieved passage supplies the \
factual or legal premise used, and the inference does not materially \
exceed the evidence.
- partial: Part of the reasoning is supported, but the AI adds an \
unsupported step, condition or conclusion.
- no: Not supported or contradicted. The relied-upon authority was \
not retrieved, the relevant proposition is absent, or the reasoning \
contradicts the passage.
- unclear: Retrieval evidence is missing, truncated or too ambiguous.
- n/a: No substantive application or reasoning claim is attached.

Respond with valid JSON only:
{
  "reasoning_grounding": "yes|partial|no|unclear|n/a",
  "retrieved_evidence": "which retrieved passage(s) are relevant, if any",
  "grounding_evidence": "2-3 sentence explanation",
  "confidence": "high|medium|low"
}
"""


BOT_CHECK_MARKERS = [
    "making sure you",
    "not a bot",
    "anubis",
    "captcha",
    "challenge",
    "cloudflare",
    "please enable javascript",
]


def fetch_judgment_text(url):
    cache_key = re.sub(r"[^a-zA-Z0-9]", "_", url)[:120]
    cache_path = CACHE_DIR / f"{cache_key}.txt"

    if cache_path.exists():
        text = cache_path.read_text(encoding="utf-8")
        if not text.startswith("[BLOCKED"):
            return text

    try:
        response = requests.get(
            url,
            timeout=30,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/128.0.0.0 Safari/537.36"
                ),
                "Accept": (
                    "text/html,application/xhtml+xml,"
                    "application/xml;q=0.9"
                ),
            },
        )
        response.raise_for_status()
    except requests.RequestException as exc:
        return f"[BLOCKED: fetch error: {exc}]"

    html = response.text
    html_lower = html.lower()

    if any(marker in html_lower for marker in BOT_CHECK_MARKERS):
        result = "[BLOCKED: bot protection detected]"
        CACHE_DIR.mkdir(parents=True, exist_ok=True)
        cache_path.write_text(result, encoding="utf-8")
        return result

    text = re.sub(
        r"<style[^>]*>.*?</style>", "", html, flags=re.DOTALL
    )
    text = re.sub(
        r"<script[^>]*>.*?</script>", "", text, flags=re.DOTALL
    )
    text = re.sub(r"<[^>]+>", "\n", text)
    text = re.sub(r"&nbsp;", " ", text)
    text = re.sub(r"&amp;", "&", text)
    text = re.sub(r"&lt;", "<", text)
    text = re.sub(r"&gt;", ">", text)
    text = re.sub(r"&#\d+;", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = text.strip()

    if len(text) < 200:
        return "[BLOCKED: response too short to be a judgment]"

    max_chars = 60_000
    if len(text) > max_chars:
        text = text[:max_chars] + "\n\n[...TRUNCATED...]"

    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    cache_path.write_text(text, encoding="utf-8")

    return text


def get_retrieved_evidence(evidence_rows, search_id):
    passages = []

    for row in evidence_rows:
        if row[0] != search_id:
            continue

        rank = row[1]
        title = str(row[3] or "")
        chunk_type = str(row[8] or "")
        chunk_summary = str(row[10] or "")
        chunk_content = str(row[11] or "")

        passages.append(
            f"[Rank {rank}] {title}\n"
            f"Chunk type: {chunk_type}\n"
            f"Summary: {chunk_summary}\n"
            f"Content: {chunk_content}"
        )

    combined = "\n\n---\n\n".join(passages)

    max_chars = 50_000
    if len(combined) > max_chars:
        combined = combined[:max_chars] + "\n\n[...TRUNCATED...]"

    return combined if combined else "[NO RETRIEVED EVIDENCE]"


def to_excel_str(value):
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def call_llm(client, system_prompt, user_prompt):
    for attempt in range(3):
        try:
            response = client.chat.completions.create(
                model=MODEL,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt},
                ],
                temperature=0.1,
                response_format={"type": "json_object"},
            )

            content = response.choices[0].message.content
            result = json.loads(content)

            return {
                key: to_excel_str(val)
                for key, val in result.items()
            }

        except Exception as exc:
            if attempt < 2:
                wait = 2 ** (attempt + 1)
                print(f"    Retry {attempt + 1} after error: {exc}")
                time.sleep(wait)
            else:
                print(f"    FAILED after 3 attempts: {exc}")
                return None

    return None


def assess_holding(client, citation, principle, judgment_text):
    source_available = not judgment_text.startswith("[BLOCKED")

    if source_available:
        system_prompt = HOLDING_SYSTEM_PROMPT_WITH_SOURCE
        user_prompt = (
            f"CITATION: {citation}\n\n"
            f"GENERATED PRINCIPLE: {principle}\n\n"
            f"SOURCE JUDGMENT TEXT:\n{judgment_text}"
        )
    else:
        system_prompt = HOLDING_SYSTEM_PROMPT_NO_SOURCE
        user_prompt = (
            f"CITATION: {citation}\n\n"
            f"GENERATED PRINCIPLE: {principle}"
        )

    result = call_llm(client, system_prompt, user_prompt)

    if result and not source_available:
        result["holding_evidence"] = (
            "[NO SOURCE TEXT - based on LLM knowledge] "
            + result.get("holding_evidence", "")
        )

    return result, source_available


def assess_grounding(client, citation, application, retrieved_text):
    user_prompt = (
        f"CITATION: {citation}\n\n"
        f"GENERATED APPLICATION: {application}\n\n"
        f"RETRIEVED PASSAGES:\n{retrieved_text}"
    )

    return call_llm(client, GROUNDING_SYSTEM_PROMPT, user_prompt)


def main():
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Print what would be done without calling the API",
    )

    parser.add_argument(
        "--limit",
        type=int,
        default=0,
        help="Process only the first N sampled units",
    )

    args = parser.parse_args()

    api_key = os.environ.get("OPENAI_API_KEY", "")

    if not api_key and not args.dry_run:
        env_path = Path(
            "F:/Projects/maternal-wellbeing-app/.env.local"
        )

        if env_path.exists():
            for line in env_path.read_text().splitlines():
                if line.startswith("OPENAI_API_KEY="):
                    api_key = line.split("=", 1)[1].strip()
                    break

    if not api_key and not args.dry_run:
        print("ERROR: No OPENAI_API_KEY found.")
        sys.exit(1)

    client = OpenAI(api_key=api_key) if not args.dry_run else None

    print(f"Workbook: {WORKBOOK_PATH}")
    print(f"Model: {MODEL}")
    print(f"Dry run: {args.dry_run}")

    wb = load_workbook(WORKBOOK_PATH)
    ws = wb["Annotations"]
    evidence_ws = wb["Retrieved Evidence"]

    # Pre-load all evidence rows for faster lookup
    print("Loading retrieved evidence...")
    evidence_rows = list(
        evidence_ws.iter_rows(
            min_row=2,
            max_row=evidence_ws.max_row,
            values_only=True,
        )
    )
    print(f"  {len(evidence_rows)} evidence rows loaded")

    timestamp = datetime.now(timezone.utc).strftime(
        "%Y-%m-%d %H:%M UTC"
    )

    processed = 0
    skipped = 0
    errors = 0
    source_fetched = 0
    source_blocked = 0

    for row_idx in range(2, ws.max_row + 1):
        in_sample = str(
            ws.cell(row_idx, 2).value or ""
        ).strip().lower()

        if in_sample != "yes":
            continue

        existing_label = str(
            ws.cell(row_idx, 22).value or ""
        ).strip()

        if existing_label:
            skipped += 1
            continue

        if args.limit and processed >= args.limit:
            break

        citation = str(ws.cell(row_idx, 7).value or "")
        principle = str(ws.cell(row_idx, 9).value or "")
        application = str(ws.cell(row_idx, 10).value or "")
        source_url = str(ws.cell(row_idx, 17).value or "")
        search_id = str(ws.cell(row_idx, 18).value or "")
        sample_order = ws.cell(row_idx, 1).value

        print(
            f"\n[{processed + 1}] Row {row_idx} | "
            f"Sample #{sample_order} | {citation[:60]}"
        )

        if args.dry_run:
            print(f"  URL: {source_url}")
            print(f"  Principle: {principle[:80]}...")
            print(f"  Application: {application[:80]}...")
            processed += 1
            continue

        # Fetch judgment text (may be blocked)
        print("  Fetching judgment...")
        judgment_text = fetch_judgment_text(source_url)
        blocked = judgment_text.startswith("[BLOCKED")

        if blocked:
            print(f"  {judgment_text}")
            source_blocked += 1
        else:
            print(
                f"  Fetched {len(judgment_text)} chars"
            )
            source_fetched += 1

        # Assess holding accuracy
        print("  Assessing holding accuracy...")
        holding_result, had_source = assess_holding(
            client, citation, principle, judgment_text
        )

        if holding_result:
            ws.cell(row_idx, 21).value = "LLM-draft"
            ws.cell(row_idx, 22).value = holding_result.get(
                "holding_accuracy", "unclear"
            )
            ws.cell(row_idx, 23).value = holding_result.get(
                "source_paragraphs", ""
            )
            ws.cell(row_idx, 24).value = holding_result.get(
                "holding_evidence", ""
            )

            print(
                f"  Holding: "
                f"{holding_result.get('holding_accuracy')}"
                f" ({holding_result.get('confidence')})"
                f"{'' if had_source else ' [no source]'}"
            )
        else:
            ws.cell(row_idx, 21).value = "LLM-draft"
            ws.cell(row_idx, 22).value = "unclear"
            ws.cell(row_idx, 24).value = "[LLM call failed]"
            errors += 1

        # Get retrieved evidence and assess grounding
        print("  Assessing reasoning grounding...")
        retrieved_text = get_retrieved_evidence(
            evidence_rows, search_id
        )

        grounding_result = assess_grounding(
            client, citation, application, retrieved_text
        )

        if grounding_result:
            ws.cell(row_idx, 25).value = grounding_result.get(
                "reasoning_grounding", "unclear"
            )
            ws.cell(row_idx, 26).value = grounding_result.get(
                "retrieved_evidence", ""
            )
            ws.cell(row_idx, 27).value = grounding_result.get(
                "grounding_evidence", ""
            )

            holding_conf = (
                holding_result.get("confidence", "")
                if holding_result
                else ""
            )
            grounding_conf = grounding_result.get(
                "confidence", ""
            )
            ws.cell(row_idx, 28).value = (
                f"holding:{holding_conf}; "
                f"grounding:{grounding_conf}"
            )

            print(
                f"  Grounding: "
                f"{grounding_result.get('reasoning_grounding')}"
                f" ({grounding_result.get('confidence')})"
            )
        else:
            ws.cell(row_idx, 25).value = "unclear"
            ws.cell(row_idx, 27).value = "[LLM call failed]"
            errors += 1

        ws.cell(row_idx, 29).value = timestamp
        processed += 1

        if processed % 10 == 0:
            print(f"\n  --- Saving workbook ({processed} done) ---")
            wb.save(WORKBOOK_PATH)

        time.sleep(0.5)

    if not args.dry_run and processed > 0:
        wb.save(WORKBOOK_PATH)

    print("\n" + "=" * 50)
    print("Pre-fill complete")
    print(f"Processed: {processed}")
    print(f"Skipped (already labelled): {skipped}")
    print(f"Source text fetched: {source_fetched}")
    print(f"Source text blocked: {source_blocked}")
    print(f"Errors: {errors}")

    if not args.dry_run and processed > 0:
        print(f"Workbook saved: {WORKBOOK_PATH}")


if __name__ == "__main__":
    main()
