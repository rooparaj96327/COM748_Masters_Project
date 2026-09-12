#!/usr/bin/env python3
"""Hercules citation, holding-accuracy, and reasoning-grounding verifier.

This program implements a reproducible screening pipeline for the COM748
Hercules audit. It reads the ``Annotations`` and ``Retrieved Evidence`` sheets
from the project workbook, predicts three labels for every citation occurrence,
and—when adjudicated human labels are present—writes validation metrics.

Important methodological boundary
---------------------------------
Holding accuracy is a legal-interpretation task. The program only predicts a
substantive holding label when authoritative judgment text is supplied through
``--authorities`` (or full source text is present in ``--registry``). Otherwise
it returns ``unclear`` rather than treating a missing source as an inaccurate
holding. Automated labels are screening results, not replacements for expert
legal review.

Example
-------
python hercules_verifier.py \
  --workbook holding_grounding_annotation.xlsx \
  --output verifier_predictions.csv \
  --metrics verifier_metrics.json

Optional authoritative-source CSV columns:
    citation, title, source_url, source_paragraphs, authoritative_text

Optional corpus-registry CSV columns (flexible aliases are accepted):
    case_id, title, citation, full_text, source_url
"""

from __future__ import annotations

import argparse
import csv
csv.field_size_limit(50 * 1024 * 1024)
import json
import math
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from difflib import SequenceMatcher
from pathlib import Path
from statistics import mean, median
from typing import Any, Iterable, Mapping, Sequence

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - exercised only without dependency
    raise SystemExit(
        "openpyxl is required. Install it with: python -m pip install openpyxl"
    ) from exc


VERSION = "1.0.0"

LABELS = ("yes", "partial", "no", "unclear", "n/a")
NON_CASE_KINDS = {"legislation", "rule", "statute", "practice_direction"}

OUTPUT_FIELDS = [
    "draft_id",
    "search_id",
    "citation",
    "citation_kind",
    "predicted_citation_status",
    "predicted_holding_accuracy",
    "predicted_reasoning_grounding",
    "holding_score",
    "grounding_score",
    "holding_source",
    "cited_authority_retrieved",
    "best_passage_rank",
    "best_passage_case_id",
    "best_passage_chunk_id",
    "best_passage_chunk_type",
    "holding_evidence",
    "grounding_evidence",
    "notes",
    "case_doc_id",
    "input_type",
    "citation_index",
    "in_validation_sample",
    "citation_verification_status",
    "adjudicated_holding_accuracy",
    "adjudicated_reasoning_grounding",
    "holding_accuracy_R1",
    "holding_accuracy_R2",
    "reasoning_grounding_R1",
    "reasoning_grounding_R2",
]

STOPWORDS = {
    "a",
    "an",
    "and",
    "are",
    "as",
    "at",
    "be",
    "been",
    "being",
    "by",
    "case",
    "court",
    "did",
    "do",
    "does",
    "for",
    "from",
    "had",
    "has",
    "have",
    "held",
    "holding",
    "in",
    "is",
    "it",
    "its",
    "law",
    "of",
    "on",
    "or",
    "principle",
    "that",
    "the",
    "their",
    "this",
    "to",
    "under",
    "was",
    "were",
    "which",
    "with",
}

NEGATION_PHRASES = (
    "no duty",
    "owed no duty",
    "does not require",
    "did not require",
    "not liable",
    "cannot establish",
    "does not establish",
    "was not entitled",
)

NEUTRAL_CITATION_RE = re.compile(
    r"\[(?P<year>\d{4})\]\s*"
    r"(?P<court>UKSC|UKHL|UKPC|EWCA\s+Civ|EWCA\s+Crim|EWHC(?:\s*\([A-Za-z]+\))?|"
    r"UKUT|UKFTT|EWFC|EWCOP|CSIH|CSOH)\s*(?P<number>\d+)",
    flags=re.IGNORECASE,
)
REPORT_CITATION_RE = re.compile(
    r"\[(?P<year>\d{4})\]\s*\d+\s+[A-Z][A-Z. ]{1,12}\s+\d+",
    flags=re.IGNORECASE,
)
LEGISLATION_RE = re.compile(
    r"\b(?:Act|Regulations?|Order|Directive|Convention|Code)\s+(?:of\s+)?\d{4}\b|"
    r"\b(?:section|s\.|article|regulation)\s*\d+[A-Za-z0-9()]*",
    flags=re.IGNORECASE,
)
RULE_RE = re.compile(
    r"\b(?:CPR|Civil Procedure Rules?|Practice Direction|PD)\b",
    flags=re.IGNORECASE,
)


@dataclass(frozen=True)
class Thresholds:
    holding_yes: float = 0.55
    holding_partial: float = 0.25
    grounding_yes: float = 0.35
    grounding_partial: float = 0.15


def clean(value: Any) -> str | None:
    """Return a stripped string or ``None`` for an empty value."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def truthy(value: Any) -> bool:
    return (clean(value) or "").lower() in {"1", "true", "yes", "y"}


def normalize_text(value: Any) -> str:
    text = (clean(value) or "").lower()
    text = text.replace("&", " and ")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def tokenize(value: Any) -> list[str]:
    return [
        token
        for token in normalize_text(value).split()
        if len(token) > 1 and token not in STOPWORDS
    ]


def citation_signatures(value: Any) -> set[str]:
    text = clean(value) or ""
    signatures: set[str] = set()
    for match in NEUTRAL_CITATION_RE.finditer(text):
        court = re.sub(r"\s+", "", match.group("court")).upper()
        signatures.add(f"{match.group('year')}:{court}:{match.group('number')}")
    for match in REPORT_CITATION_RE.finditer(text):
        signatures.add(normalize_text(match.group(0)))
    return signatures


def classify_citation(citation: Any, annotation_type: Any = None) -> str:
    text = clean(citation) or ""
    annotation = (clean(annotation_type) or "").lower()
    if annotation in {"legislation", "statute", "statutory_provision"} or LEGISLATION_RE.search(text):
        return "legislation"
    if annotation in {"rule", "practice_direction"} or RULE_RE.search(text):
        return "rule"
    if NEUTRAL_CITATION_RE.search(text):
        return "neutral"
    if REPORT_CITATION_RE.search(text):
        return "report"
    return "unknown"


def party_tokens(value: Any) -> set[str]:
    """Extract stable party-name tokens while dropping citation material."""
    text = clean(value) or ""
    text = NEUTRAL_CITATION_RE.sub(" ", text)
    text = REPORT_CITATION_RE.sub(" ", text)
    tokens = {
        token
        for token in tokenize(text)
        if token
        not in {
            "ltd",
            "limited",
            "plc",
            "another",
            "anor",
            "others",
            "ors",
            "appellant",
            "respondent",
            "secretary",
            "state",
        }
    }
    return tokens


def authority_match_score(query: Any, candidate: Any) -> float:
    query_text = clean(query) or ""
    candidate_text = clean(candidate) or ""
    if not query_text or not candidate_text:
        return 0.0

    query_sigs = citation_signatures(query_text)
    candidate_sigs = citation_signatures(candidate_text)
    if query_sigs and candidate_sigs and query_sigs.intersection(candidate_sigs):
        return 1.0

    left = party_tokens(query_text)
    right = party_tokens(candidate_text)
    if not left or not right:
        return 0.0
    intersection = len(left.intersection(right))
    union = len(left.union(right))
    jaccard = intersection / union if union else 0.0
    containment = intersection / min(len(left), len(right))
    sequence = SequenceMatcher(None, normalize_text(query_text), normalize_text(candidate_text)).ratio()
    return max(jaccard, 0.75 * containment + 0.25 * sequence)


def split_passages(text: Any, max_words: int = 120) -> list[str]:
    raw = clean(text) or ""
    if not raw:
        return []
    rough = [piece.strip() for piece in re.split(r"\n+|(?<=[.!?])\s+", raw) if piece.strip()]
    chunks: list[str] = []
    current: list[str] = []
    for piece in rough:
        words = piece.split()
        if current and len(current) + len(words) > max_words:
            chunks.append(" ".join(current))
            current = []
        if len(words) > max_words:
            for start in range(0, len(words), max_words):
                chunks.append(" ".join(words[start : start + max_words]))
        else:
            current.extend(words)
    if current:
        chunks.append(" ".join(current))
    return chunks or [raw]


def _tfidf_vectors(documents: Sequence[str]) -> list[dict[str, float]]:
    token_counts = [Counter(tokenize(document)) for document in documents]
    total_docs = len(token_counts)
    document_frequency = Counter()
    for counts in token_counts:
        document_frequency.update(counts.keys())

    vectors: list[dict[str, float]] = []
    for counts in token_counts:
        total_terms = sum(counts.values()) or 1
        vector = {
            token: (count / total_terms) * (math.log((1 + total_docs) / (1 + document_frequency[token])) + 1)
            for token, count in counts.items()
        }
        vectors.append(vector)
    return vectors


def _cosine(left: Mapping[str, float], right: Mapping[str, float]) -> float:
    common = set(left).intersection(right)
    numerator = sum(left[token] * right[token] for token in common)
    left_norm = math.sqrt(sum(value * value for value in left.values()))
    right_norm = math.sqrt(sum(value * value for value in right.values()))
    if not left_norm or not right_norm:
        return 0.0
    return numerator / (left_norm * right_norm)


def passage_similarity(claim: Any, passages: Sequence[str]) -> tuple[float, str | None]:
    claim_text = clean(claim) or ""
    usable = [clean(passage) or "" for passage in passages if clean(passage)]
    if not claim_text or not usable:
        return 0.0, None

    vectors = _tfidf_vectors([claim_text, *usable])
    claim_tokens = set(tokenize(claim_text))
    best_score = 0.0
    best_passage: str | None = None
    for index, passage in enumerate(usable, start=1):
        passage_tokens = set(tokenize(passage))
        cosine = _cosine(vectors[0], vectors[index])
        coverage = (
            len(claim_tokens.intersection(passage_tokens)) / len(claim_tokens)
            if claim_tokens
            else 0.0
        )
        sequence = SequenceMatcher(
            None, normalize_text(claim_text)[:600], normalize_text(passage)[:1200]
        ).ratio()
        score = min(1.0, 0.60 * cosine + 0.35 * coverage + 0.05 * sequence)
        if score > best_score:
            best_score = score
            best_passage = passage
    return best_score, best_passage


def polarity_mismatch(claim: Any, passage: Any) -> bool:
    claim_text = normalize_text(claim)
    passage_text = normalize_text(passage)
    if not claim_text or not passage_text:
        return False
    claim_negative = any(phrase in claim_text for phrase in NEGATION_PHRASES)
    passage_negative = any(phrase in passage_text for phrase in NEGATION_PHRASES)
    shared_legal_terms = set(tokenize(claim_text)).intersection(tokenize(passage_text)).intersection(
        {"duty", "liable", "liability", "require", "required", "entitled", "permit", "prohibit"}
    )
    return bool(shared_legal_terms) and claim_negative != passage_negative


def read_sheet_records(sheet: Any, header_row: int = 1) -> list[dict[str, Any]]:
    headers = [clean(cell.value) for cell in sheet[header_row]]
    records: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(clean(value) is not None for value in values):
            continue
        records.append(
            {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header is not None
            }
        )
    return records


def require_columns(records: Sequence[Mapping[str, Any]], required: Iterable[str], source: str) -> None:
    if not records:
        raise ValueError(f"{source} contains no data rows")
    available = set(records[0])
    missing = sorted(set(required) - available)
    if missing:
        raise ValueError(f"{source} is missing required columns: {', '.join(missing)}")


def read_csv_records(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as stream:
        return list(csv.DictReader(stream))


def flexible_value(record: Mapping[str, Any], names: Sequence[str]) -> Any:
    for name in names:
        if clean(record.get(name)) is not None:
            return record.get(name)
    return None


def standardize_source_record(record: Mapping[str, Any]) -> dict[str, Any]:
    return {
        "case_id": flexible_value(record, ("case_id", "id", "retrieved_case_id", "corpus_case_id")),
        "title": flexible_value(record, ("title", "case_title", "retrieved_case_title")),
        "citation": flexible_value(record, ("citation", "neutral_citation", "case_citation")),
        "text": flexible_value(
            record,
            ("authoritative_text", "full_text", "judgment_text", "text", "chunk_content"),
        ),
        "source_url": flexible_value(record, ("source_url", "url", "authoritative_source_url")),
        "source_paragraphs": flexible_value(record, ("source_paragraphs", "paragraphs", "section")),
    }


def source_description(source: Mapping[str, Any]) -> str:
    return " ".join(
        part
        for part in [clean(source.get("title")), clean(source.get("citation"))]
        if part
    )


def find_source(citation: Any, sources: Sequence[Mapping[str, Any]], threshold: float = 0.58) -> tuple[dict[str, Any] | None, float]:
    best: dict[str, Any] | None = None
    best_score = 0.0
    for source in sources:
        score = authority_match_score(citation, source_description(source))
        if score > best_score:
            best = dict(source)
            best_score = score
    return (best, best_score) if best is not None and best_score >= threshold else (None, best_score)


def evidence_description(record: Mapping[str, Any]) -> str:
    return " ".join(
        part
        for part in [
            clean(record.get("retrieved_case_title")),
            clean(record.get("cited_cases")),
        ]
        if part
    )


def find_retrieved_authority(citation: Any, evidence: Sequence[Mapping[str, Any]]) -> tuple[dict[str, Any] | None, float]:
    best: dict[str, Any] | None = None
    best_score = 0.0
    for record in evidence:
        score = authority_match_score(citation, evidence_description(record))
        if score > best_score:
            best = dict(record)
            best_score = score
    return (best, best_score) if best is not None and best_score >= 0.58 else (None, best_score)


def best_grounding_passage(
    claim: Any, evidence: Sequence[Mapping[str, Any]]
) -> tuple[float, dict[str, Any] | None, str | None]:
    best_score = 0.0
    best_record: dict[str, Any] | None = None
    best_text: str | None = None
    for record in evidence:
        text = clean(record.get("chunk_content")) or clean(record.get("chunk_summary"))
        if not text:
            continue
        score, passage = passage_similarity(claim, split_passages(text))
        if score > best_score:
            best_score = score
            best_record = dict(record)
            best_text = passage
    return best_score, best_record, best_text


def predict_citation_status(
    annotation: Mapping[str, Any],
    citation_kind: str,
    retrieved_authority: Mapping[str, Any] | None,
    registry_source: Mapping[str, Any] | None,
    use_workbook_flags: bool,
) -> tuple[str, str]:
    if citation_kind in NON_CASE_KINDS:
        return "not_applicable_non_case_reference", "non-case reference"
    if retrieved_authority is not None:
        return "retrieved_and_cited", "citation matched an authority in the linked retrieval"

    linked_flag = (clean(annotation.get("citation_in_linked_retrieval")) or "").lower()
    if use_workbook_flags and linked_flag == "yes":
        return (
            "retrieved_id_citation_mismatch",
            "retrieval linkage was flagged present, but the retrieved title/citation did not match",
        )

    registry_flag = (clean(annotation.get("citation_in_corpus_registry")) or "").lower()
    if registry_source is not None or (use_workbook_flags and registry_flag == "yes"):
        return "in_corpus_not_retrieved", "authority was found in the corpus but not in the linked retrieval"
    return (
        "not_found_in_corpus_registry",
        "authority was not found in the supplied corpus registry",
    )


def predict_holding(
    annotation: Mapping[str, Any],
    citation_kind: str,
    authoritative_source: Mapping[str, Any] | None,
    registry_source: Mapping[str, Any] | None,
    thresholds: Thresholds,
) -> tuple[str, float, str, str]:
    if citation_kind in NON_CASE_KINDS:
        return "n/a", 0.0, "none", "holding accuracy is not applicable to this non-case reference"

    source = authoritative_source
    source_kind = "authoritative"
    if source is None or clean(source.get("text")) is None:
        source = registry_source
        source_kind = "corpus"
    if source is None or clean(source.get("text")) is None:
        return (
            "unclear",
            0.0,
            "none",
            "authoritative judgment text unavailable; holding cannot be verified from a citation alone",
        )

    principle = clean(annotation.get("generated_principle")) or ""
    score, passage = passage_similarity(principle, split_passages(source.get("text")))
    if not principle or passage is None:
        return "unclear", score, source_kind, "generated principle or usable source text is missing"

    mismatch = polarity_mismatch(principle, passage)
    if mismatch:
        label = "no"
    elif score >= thresholds.holding_yes:
        label = "yes"
    elif score >= thresholds.holding_partial:
        label = "partial"
    else:
        label = "no"

    excerpt = re.sub(r"\s+", " ", passage).strip()[:360]
    evidence = f"best {source_kind} match (score={score:.3f}): {excerpt}"
    if mismatch:
        evidence += " [possible polarity/negation conflict]"
    return label, score, source_kind, evidence


def predict_grounding(
    annotation: Mapping[str, Any],
    citation_kind: str,
    evidence: Sequence[Mapping[str, Any]],
    retrieved_authority: Mapping[str, Any] | None,
    thresholds: Thresholds,
) -> tuple[str, float, dict[str, Any] | None, str]:
    application = clean(annotation.get("generated_application"))
    if not application:
        return "n/a", 0.0, None, "no substantive generated application was supplied"
    if not clean(annotation.get("search_id")) or not evidence:
        return "unclear", 0.0, None, "linked retrieval evidence is missing"

    score, best_record, best_passage = best_grounding_passage(application, evidence)
    if best_record is None or best_passage is None:
        return "unclear", score, None, "retrieval rows exist but contain no usable passage text"

    case_reference = citation_kind not in NON_CASE_KINDS
    authority_present = retrieved_authority is not None
    if case_reference and not authority_present:
        label = "no"
    elif score >= thresholds.grounding_yes:
        label = "yes"
    elif score >= thresholds.grounding_partial:
        label = "partial"
    else:
        label = "no"

    excerpt = re.sub(r"\s+", " ", best_passage).strip()[:360]
    rank = clean(best_record.get("retrieval_rank")) or "?"
    chunk_type = clean(best_record.get("best_chunk_type")) or "unknown"
    if case_reference and not authority_present:
        evidence_note = (
            f"cited authority absent from the linked retrieval; closest passage rank {rank} "
            f"({chunk_type}, score={score:.3f}): {excerpt}"
        )
    else:
        evidence_note = f"best linked passage rank {rank} ({chunk_type}, score={score:.3f}): {excerpt}"
    return label, score, best_record, evidence_note


def distribution(records: Sequence[Mapping[str, Any]], field: str) -> dict[str, int]:
    return dict(sorted(Counter(clean(record.get(field)) for record in records if clean(record.get(field))).items()))


def classification_metrics(
    records: Sequence[Mapping[str, Any]], prediction_field: str, truth_field: str
) -> dict[str, Any]:
    pairs = [
        (clean(record.get(prediction_field)), clean(record.get(truth_field)))
        for record in records
        if clean(record.get(prediction_field)) and clean(record.get(truth_field))
    ]
    pairs = [(prediction, truth) for prediction, truth in pairs if prediction and truth]
    labels = sorted({value for pair in pairs for value in pair})
    count = len(pairs)
    if not count:
        return {"paired_rows": 0}

    matrix = {prediction: {truth: 0 for truth in labels} for prediction in labels}
    for prediction, truth in pairs:
        matrix[prediction][truth] += 1

    observed = sum(prediction == truth for prediction, truth in pairs) / count
    pred_counts = Counter(prediction for prediction, _ in pairs)
    truth_counts = Counter(truth for _, truth in pairs)
    expected = sum(
        (pred_counts[label] / count) * (truth_counts[label] / count) for label in labels
    )
    kappa = (observed - expected) / (1 - expected) if expected != 1 else None

    per_class: dict[str, dict[str, float | int | None]] = {}
    recalls: list[float] = []
    f1_values: list[float] = []
    for label in labels:
        tp = sum(prediction == label and truth == label for prediction, truth in pairs)
        fp = sum(prediction == label and truth != label for prediction, truth in pairs)
        fn = sum(prediction != label and truth == label for prediction, truth in pairs)
        precision = tp / (tp + fp) if tp + fp else None
        recall = tp / (tp + fn) if tp + fn else None
        f1 = (
            2 * precision * recall / (precision + recall)
            if precision is not None and recall is not None and precision + recall
            else None
        )
        if recall is not None:
            recalls.append(recall)
        if f1 is not None:
            f1_values.append(f1)
        per_class[label] = {
            "support": truth_counts[label],
            "precision": precision,
            "recall": recall,
            "f1": f1,
        }

    return {
        "paired_rows": count,
        "exact_accuracy": observed,
        "balanced_accuracy": mean(recalls) if recalls else None,
        "cohens_kappa": kappa,
        "macro_f1_defined_classes": mean(f1_values) if f1_values else None,
        "confusion_matrix_predicted_by_truth": matrix,
        "per_class": per_class,
    }


def csv_value(value: Any) -> Any:
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return "" if value is None else value


def build_prediction(
    annotation: Mapping[str, Any],
    evidence: Sequence[Mapping[str, Any]],
    registry: Sequence[Mapping[str, Any]],
    authorities: Sequence[Mapping[str, Any]],
    thresholds: Thresholds,
    use_workbook_flags: bool,
) -> dict[str, Any]:
    citation = clean(annotation.get("citation")) or ""
    citation_kind = classify_citation(citation, annotation.get("citation_type"))
    retrieved_authority, _ = find_retrieved_authority(citation, evidence)
    registry_source, _ = find_source(citation, registry)
    authoritative_source, _ = find_source(citation, authorities)

    citation_status, citation_note = predict_citation_status(
        annotation,
        citation_kind,
        retrieved_authority,
        registry_source,
        use_workbook_flags,
    )
    holding_label, holding_score, holding_source, holding_evidence = predict_holding(
        annotation,
        citation_kind,
        authoritative_source,
        registry_source,
        thresholds,
    )
    grounding_label, grounding_score, best_record, grounding_evidence = predict_grounding(
        annotation,
        citation_kind,
        evidence,
        retrieved_authority,
        thresholds,
    )

    best_record = best_record or {}
    output = {
        "draft_id": annotation.get("draft_id"),
        "search_id": annotation.get("search_id"),
        "citation": annotation.get("citation"),
        "citation_kind": citation_kind,
        "predicted_citation_status": citation_status,
        "predicted_holding_accuracy": holding_label,
        "predicted_reasoning_grounding": grounding_label,
        "holding_score": f"{holding_score:.4f}",
        "grounding_score": f"{grounding_score:.4f}",
        "holding_source": holding_source,
        "cited_authority_retrieved": retrieved_authority is not None,
        "best_passage_rank": best_record.get("retrieval_rank"),
        "best_passage_case_id": best_record.get("retrieved_case_id"),
        "best_passage_chunk_id": best_record.get("best_chunk_id"),
        "best_passage_chunk_type": best_record.get("best_chunk_type"),
        "holding_evidence": holding_evidence,
        "grounding_evidence": grounding_evidence,
        "notes": citation_note,
        "case_doc_id": annotation.get("case_doc_id"),
        "input_type": annotation.get("input_type"),
        "citation_index": annotation.get("citation_index"),
        "in_validation_sample": annotation.get("in_validation_sample"),
        "citation_verification_status": annotation.get("citation_verification_status"),
        "adjudicated_holding_accuracy": annotation.get("adjudicated_holding_accuracy"),
        "adjudicated_reasoning_grounding": annotation.get("adjudicated_reasoning_grounding"),
        "holding_accuracy_R1": annotation.get("holding_accuracy_R1"),
        "holding_accuracy_R2": annotation.get("holding_accuracy_R2"),
        "reasoning_grounding_R1": annotation.get("reasoning_grounding_R1"),
        "reasoning_grounding_R2": annotation.get("reasoning_grounding_R2"),
    }
    return {field: csv_value(output.get(field)) for field in OUTPUT_FIELDS}


def verify(args: argparse.Namespace) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    workbook_path = Path(args.workbook)
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if args.annotations_sheet not in workbook.sheetnames:
        raise ValueError(f"Workbook does not contain sheet: {args.annotations_sheet}")
    if args.evidence_sheet not in workbook.sheetnames:
        raise ValueError(f"Workbook does not contain sheet: {args.evidence_sheet}")

    annotations = read_sheet_records(workbook[args.annotations_sheet])
    evidence_rows = read_sheet_records(workbook[args.evidence_sheet])
    require_columns(
        annotations,
        ("draft_id", "citation", "generated_principle", "generated_application", "search_id"),
        args.annotations_sheet,
    )
    require_columns(evidence_rows, ("search_id", "retrieval_rank"), args.evidence_sheet)

    if args.sample_only:
        annotations = [
            row for row in annotations if (clean(row.get("in_validation_sample")) or "").lower() == "yes"
        ]

    evidence_by_search: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in evidence_rows:
        search_id = clean(row.get("search_id"))
        if search_id:
            evidence_by_search[search_id].append(row)

    registry = (
        [standardize_source_record(row) for row in read_csv_records(Path(args.registry))]
        if args.registry
        else []
    )
    authorities = (
        [standardize_source_record(row) for row in read_csv_records(Path(args.authorities))]
        if args.authorities
        else []
    )

    thresholds = Thresholds(
        holding_yes=args.holding_yes,
        holding_partial=args.holding_partial,
        grounding_yes=args.grounding_yes,
        grounding_partial=args.grounding_partial,
    )
    if thresholds.holding_partial > thresholds.holding_yes:
        raise ValueError("--holding-partial cannot exceed --holding-yes")
    if thresholds.grounding_partial > thresholds.grounding_yes:
        raise ValueError("--grounding-partial cannot exceed --grounding-yes")

    predictions = []
    for annotation in annotations:
        search_id = clean(annotation.get("search_id")) or ""
        predictions.append(
            build_prediction(
                annotation,
                evidence_by_search.get(search_id, []),
                registry,
                authorities,
                thresholds,
                not args.ignore_workbook_registry_flags,
            )
        )

    validation_rows = [
        record
        for record in predictions
        if (clean(record.get("in_validation_sample")) or "").lower() == "yes"
    ]
    scores = {
        "holding": [float(record["holding_score"]) for record in predictions],
        "grounding": [float(record["grounding_score"]) for record in predictions],
    }
    metrics = {
        "verifier_version": VERSION,
        "input_workbook": workbook_path.name,
        "parameters": {
            "sample_only": args.sample_only,
            "use_workbook_registry_flags": not args.ignore_workbook_registry_flags,
            "authoritative_sources_supplied": bool(args.authorities),
            "corpus_registry_supplied": bool(args.registry),
            "thresholds": thresholds.__dict__,
        },
        "coverage": {
            "citation_occurrences": len(predictions),
            "unique_drafts": len({clean(record.get("draft_id")) for record in predictions}),
            "validation_occurrences": len(validation_rows),
        },
        "prediction_distributions": {
            "citation_status": distribution(predictions, "predicted_citation_status"),
            "holding_accuracy": distribution(predictions, "predicted_holding_accuracy"),
            "reasoning_grounding": distribution(predictions, "predicted_reasoning_grounding"),
            "holding_source": distribution(predictions, "holding_source"),
        },
        "score_summaries": {
            name: {
                "minimum": min(values) if values else None,
                "median": median(values) if values else None,
                "mean": mean(values) if values else None,
                "maximum": max(values) if values else None,
            }
            for name, values in scores.items()
        },
        "validation": {
            "citation_status": classification_metrics(
                validation_rows, "predicted_citation_status", "citation_verification_status"
            ),
            "holding_accuracy": classification_metrics(
                validation_rows, "predicted_holding_accuracy", "adjudicated_holding_accuracy"
            ),
            "reasoning_grounding": classification_metrics(
                validation_rows,
                "predicted_reasoning_grounding",
                "adjudicated_reasoning_grounding",
            ),
        },
    }
    return predictions, metrics


def write_predictions(path: Path, records: Sequence[Mapping[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=OUTPUT_FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)


def write_metrics(path: Path, metrics: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as stream:
        json.dump(metrics, stream, indent=2, ensure_ascii=False)
        stream.write("\n")


def percentage(value: Any) -> str:
    return "n/a" if value is None else f"{100 * float(value):.1f}%"


def print_summary(output_path: Path, metrics_path: Path, metrics: Mapping[str, Any]) -> None:
    coverage = metrics["coverage"]
    print(
        f"Verified {coverage['citation_occurrences']} citation occurrences "
        f"from {coverage['unique_drafts']} draft judgments."
    )
    print(f"Predictions: {output_path}")
    print(f"Metrics:     {metrics_path}")
    for label, result in metrics["validation"].items():
        if result.get("paired_rows"):
            print(
                f"{label}: n={result['paired_rows']}, "
                f"accuracy={percentage(result.get('exact_accuracy'))}, "
                f"balanced_accuracy={percentage(result.get('balanced_accuracy'))}, "
                f"kappa={result.get('cohens_kappa') if result.get('cohens_kappa') is not None else 'n/a'}"
            )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Screen Hercules citation occurrences for corpus/retrieval status, holding accuracy, "
            "and reasoning grounding; optionally validate against adjudicated labels."
        )
    )
    parser.add_argument("--version", action="version", version=f"%(prog)s {VERSION}")
    parser.add_argument("--workbook", required=True, help="Annotation workbook (.xlsx)")
    parser.add_argument("--output", required=True, help="Destination predictions CSV")
    parser.add_argument(
        "--metrics",
        help="Destination metrics JSON (default: <output-stem>_metrics.json)",
    )
    parser.add_argument(
        "--authorities",
        help="Optional CSV containing authoritative judgment text for holding screening",
    )
    parser.add_argument(
        "--registry",
        help="Optional CSV containing the corpus registry and, optionally, full case text",
    )
    parser.add_argument("--annotations-sheet", default="Annotations")
    parser.add_argument("--evidence-sheet", default="Retrieved Evidence")
    parser.add_argument("--sample-only", action="store_true", help="Process only validation-sample rows")
    parser.add_argument(
        "--ignore-workbook-registry-flags",
        action="store_true",
        help="Do not use prefilled citation-in-corpus/retrieval flags as automated input features",
    )
    parser.add_argument("--holding-yes", type=float, default=0.55)
    parser.add_argument("--holding-partial", type=float, default=0.25)
    parser.add_argument("--grounding-yes", type=float, default=0.35)
    parser.add_argument("--grounding-partial", type=float, default=0.15)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    output_path = Path(args.output)
    metrics_path = Path(args.metrics) if args.metrics else output_path.with_name(
        f"{output_path.stem}_metrics.json"
    )
    try:
        predictions, metrics = verify(args)
        write_predictions(output_path, predictions)
        write_metrics(metrics_path, metrics)
        print_summary(output_path, metrics_path, metrics)
        return 0
    except (OSError, ValueError, KeyError) as exc:
        parser.exit(2, f"error: {exc}\n")


if __name__ == "__main__":
    sys.exit(main())