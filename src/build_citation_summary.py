from collections import Counter, defaultdict
from pathlib import Path
import csv
import json

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
INPUT_FILE = (
    PROJECT_ROOT
    / "data"
    / "processed"
    / "unique_citation_review.xlsx"
)
OUTPUT_DIRECTORY = PROJECT_ROOT / "data" / "processed"


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def number(value):
    if value in (None, ""):
        return 0
    return int(value)


def main():
    if not INPUT_FILE.exists():
        raise FileNotFoundError(f"File not found: {INPUT_FILE}")

    workbook = load_workbook(INPUT_FILE, data_only=True)
    worksheet = workbook["unique_citation_review"]

    headers = [clean(cell.value) for cell in worksheet[1]]

    rows = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, values))

        if clean(record.get("authority")):
            rows.append(record)

    required_review_fields = [
        "case_exists_external",
        "citation_correct_external",
        "authoritative_source_url",
        "review_notes",
    ]

    incomplete_rows = [
        clean(record["authority"])
        for record in rows
        if any(not clean(record.get(field)) for field in required_review_fields)
    ]

    if incomplete_rows:
        raise ValueError(
            "External review is incomplete for: "
            + "; ".join(incomplete_rows)
        )

    total_mentions = sum(number(row["mention_count"]) for row in rows)
    real_mentions = sum(number(row["real_mentions"]) for row in rows)
    synthetic_mentions = sum(
        number(row["synthetic_mentions"]) for row in rows
    )

    citation_types = Counter(
        clean(row["citation_type"]).lower() for row in rows
    )
    priorities = Counter(
        clean(row["review_priority"]).lower() for row in rows
    )
    existence_results = Counter(
        clean(row["case_exists_external"]).lower() for row in rows
    )
    correctness_results = Counter(
        clean(row["citation_correct_external"]).lower() for row in rows
    )

    summary_metrics = {
        "reviewed_authority_groups": len(rows),
        "citation_mentions": total_mentions,
        "real_case_mentions": real_mentions,
        "synthetic_case_mentions": synthetic_mentions,
        "case_law_authority_groups": citation_types["case_law"],
        "non_case_reference_groups": citation_types[
            "legislation_or_rule"
        ],
        "externally_confirmed_cases": existence_results["yes"],
        "ambiguous_case_references": existence_results["ambiguous"],
        "non_case_references": existence_results["n/a"],
        "correct_citations": correctness_results["yes"],
        "incomplete_citations": correctness_results["incomplete"],
        "incorrect_citations": correctness_results["no"],
        "high_priority_reviews": priorities["high"],
        "medium_priority_reviews": priorities["medium"],
        "low_priority_reviews": priorities["low"],
    }

    breakdown = defaultdict(
        lambda: {
            "authority_groups": 0,
            "mentions": 0,
            "real_mentions": 0,
            "synthetic_mentions": 0,
        }
    )

    for row in rows:
        status = clean(
            row["citation_correct_external"]
        ).lower()

        breakdown[status]["authority_groups"] += 1
        breakdown[status]["mentions"] += number(row["mention_count"])
        breakdown[status]["real_mentions"] += number(
            row["real_mentions"]
        )
        breakdown[status]["synthetic_mentions"] += number(
            row["synthetic_mentions"]
        )

    OUTPUT_DIRECTORY.mkdir(parents=True, exist_ok=True)

    summary_csv = OUTPUT_DIRECTORY / "citation_review_summary.csv"
    summary_json = OUTPUT_DIRECTORY / "citation_review_summary.json"
    breakdown_csv = (
        OUTPUT_DIRECTORY / "citation_correctness_breakdown.csv"
    )

    with summary_csv.open("w", newline="", encoding="utf-8") as file:
        writer = csv.writer(file)
        writer.writerow(["metric", "value"])

        for metric, value in summary_metrics.items():
            writer.writerow([metric, value])

    with breakdown_csv.open(
        "w", newline="", encoding="utf-8"
    ) as file:
        writer = csv.writer(file)
        writer.writerow(
            [
                "citation_status",
                "authority_groups",
                "mentions",
                "real_mentions",
                "synthetic_mentions",
            ]
        )

        for status in ["yes", "incomplete", "no"]:
            values = breakdown[status]
            writer.writerow(
                [
                    status,
                    values["authority_groups"],
                    values["mentions"],
                    values["real_mentions"],
                    values["synthetic_mentions"],
                ]
            )

    with summary_json.open("w", encoding="utf-8") as file:
        json.dump(
            {
                "summary": summary_metrics,
                "citation_correctness_breakdown": dict(breakdown),
            },
            file,
            indent=2,
        )

    print("\nCitation review summary completed")
    print("---------------------------------")

    for metric, value in summary_metrics.items():
        print(f"{metric}: {value}")

    print(f"\nSummary CSV: {summary_csv}")
    print(f"Breakdown CSV: {breakdown_csv}")
    print(f"Summary JSON: {summary_json}")


if __name__ == "__main__":
    main()